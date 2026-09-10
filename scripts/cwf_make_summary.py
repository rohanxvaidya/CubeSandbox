#!/usr/bin/env python3
# Regenerate the usual sweep outputs from a CWF create-delete sweep dir:
#   - cwf_create_latency_vs_concurrency.png (create avg/p95 vs concurrency + SLA crossing)
#   - per_sandbox_latency.png (per-sandbox create latency distribution per concurrency)
#   - <dirname>_summary.xlsx (summary + per-point sheets with both graphs embedded)
import json, glob, os, sys, csv
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage

OUT = sys.argv[1]
SLA = float(sys.argv[2]) if len(sys.argv) > 2 else 200.0
name = os.path.basename(os.path.normpath(OUT))

points = []
for f in sorted(glob.glob(os.path.join(OUT, "c*_n*.json"))):
    j = json.load(open(f))
    cfg, cr, de, s = j["config"], j["create"], j.get("delete", {}), j["summary"]
    points.append({
        "c": cfg["concurrency"], "n": cfg["total"],
        "avg": cr["avg"], "p50": cr["p50"], "p95": cr["p95"], "p99": cr["p99"],
        "del_avg": de.get("avg", 0.0),
        "tput": s["throughput_qps"], "succ": s["success_rate"] * 100, "err": s["errors"],
        "raw": [r["create_ms"] for r in j.get("raw", [])],
    })
points.sort(key=lambda p: p["c"])
if not points:
    sys.exit("no c*_n*.json points found in %s" % OUT)

# ---- CSV (in case not already present) ----
csv_path = os.path.join(OUT, "cwf_sweep_createdelete.csv")
with open(csv_path, "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["concurrency","n_total","create_avg_ms","create_p50_ms","create_p95_ms",
                "create_p99_ms","delete_avg_ms","throughput_sb_s","success_pct","errors"])
    for p in points:
        w.writerow([p["c"],p["n"],round(p["avg"],3),round(p["p50"],3),round(p["p95"],3),
                    round(p["p99"],3),round(p["del_avg"],3),round(p["tput"],3),
                    round(p["succ"],1),p["err"]])

# ---- Graph 1: create latency vs concurrency ----
c = [p["c"] for p in points]; avg = [p["avg"] for p in points]; p95 = [p["p95"] for p in points]
knee = None
for i in range(1, len(c)):
    if (avg[i-1]-SLA)*(avg[i]-SLA) <= 0 and avg[i] != avg[i-1]:
        knee = c[i-1] + (c[i]-c[i-1])*(SLA-avg[i-1])/(avg[i]-avg[i-1]); break
g1 = os.path.join(OUT, "cwf_create_latency_vs_concurrency.png")
plt.figure(figsize=(9, 5.5))
plt.plot(c, avg, "-o", label="create avg (ms)")
plt.plot(c, p95, "-s", label="create p95 (ms)", alpha=0.7)
plt.axhline(SLA, color="red", ls="--", label=f"{SLA:.0f} ms SLA")
if knee: plt.axvline(knee, color="green", ls=":", label=f"SLA crossing ~c{knee:.0f}")
plt.xlabel("concurrency"); plt.ylabel("create latency (ms)")
plt.title("CubeSandbox create latency vs concurrency (CWF)")
plt.legend(); plt.grid(True, alpha=0.3); plt.tight_layout(); plt.savefig(g1, dpi=120); plt.close()

# ---- Graph 2: per-sandbox create latency distribution per concurrency ----
g2 = os.path.join(OUT, "per_sandbox_latency.png")
data = [p["raw"] for p in points if p["raw"]]
labels = [p["c"] for p in points if p["raw"]]
plt.figure(figsize=(max(9, len(labels)*0.5), 5.5))
if data:
    plt.boxplot(data, tick_labels=[str(l) for l in labels], showfliers=False)
plt.xlabel("concurrency"); plt.ylabel("per-sandbox create latency (ms)")
plt.title("Per-sandbox create latency distribution (CWF)")
plt.grid(True, axis="y", alpha=0.3); plt.tight_layout(); plt.savefig(g2, dpi=120); plt.close()

# ---- Excel summary ----
wb = Workbook()
ws = wb.active; ws.title = "summary"
peak = max(points, key=lambda p: p["tput"])
ws.append(["Metric", "Value"])
for k, v in [
    ("run", name),
    ("points", len(points)),
    ("concurrency range", f"{c[0]}..{c[-1]}"),
    ("SLA (ms)", SLA),
    ("SLA crossing (knee)", f"c{knee:.0f}" if knee else "none"),
    ("peak throughput (sb/s)", f"{peak['tput']:.1f} @ c{peak['c']}"),
    ("min success %", min(p["succ"] for p in points)),
    ("total errors", sum(p["err"] for p in points)),
]:
    ws.append([k, v])

wp = wb.create_sheet("per_point")
wp.append(["concurrency","n_total","create_avg_ms","create_p50_ms","create_p95_ms",
           "create_p99_ms","delete_avg_ms","throughput_sb_s","success_pct","errors"])
for p in points:
    wp.append([p["c"],p["n"],round(p["avg"],2),round(p["p50"],2),round(p["p95"],2),
               round(p["p99"],2),round(p["del_avg"],2),round(p["tput"],1),round(p["succ"],1),p["err"]])

wg = wb.create_sheet("graphs")
wg.add_image(XLImage(g1), "A1")
wg.add_image(XLImage(g2), "A32")
xlsx = os.path.join(OUT, f"{name}_summary.xlsx")
wb.save(xlsx)
print(f"summary: {xlsx}")
print(f"graphs: {g1} ; {g2}")
print(f"peak throughput {peak['tput']:.1f} sb/s @ c{peak['c']}; knee " + (f"c{knee:.0f}" if knee else "none"))
